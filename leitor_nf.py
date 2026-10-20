"""Leitor de NFs (DANFe/NFS-e em PDF ou XML) -> planilha Excel.

Extrai: paciente, número da NF, valor bruto, valor líquido e mês/ano de referência.
Uso: dê dois cliques (abre a janela) ou `python leitor_nf.py arquivo1.pdf pasta ...`
"""
import os
import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass, asdict

MESES = ["JANEIRO", "FEVEREIRO", "MARCO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
VALOR = r"(?:R\$\s*)?(\d{1,3}(?:\.\d{3})*,\d{2}|\d+,\d{2})"


@dataclass
class Nota:
    arquivo: str = ""
    paciente: str = ""
    numero: str = ""
    valor_bruto: float | None = None
    valor_liquido: float | None = None
    mes_ano: str = ""
    periodo: str = ""
    aviso: str = ""


# ---------------------------------------------------------------- utilidades
def sem_acento(s: str) -> str:
    tabela = str.maketrans("ÁÀÂÃÄÉÈÊËÍÌÎÏÓÒÔÕÖÚÙÛÜÇáàâãäéèêëíìîïóòôõöúùûüç",
                           "AAAAAEEEEIIIIOOOOOUUUUCaaaaaeeeeiiiiooooouuuuc")
    return s.translate(tabela)


def num_br(txt: str) -> float | None:
    try:
        return float(txt.replace(".", "").replace(",", "."))
    except (ValueError, AttributeError):
        return None


def extrair_periodo(texto: str) -> tuple[str, str]:
    """Retorna (MM/AAAA, trecho do período) a partir da descrição."""
    t = sem_acento(texto).upper()
    t = re.sub(r"\s+", " ", t)
    meses = "|".join(MESES)
    # "PERIODO DE 01 A 28 DE FEVEREIRO DE 2026"  /  "DE 25 DE JANEIRO A 10 DE FEVEREIRO DE 2026"
    m = re.search(rf"PERIODO\s+(?:DE\s+)?(\d{{1,2}})(?:\s+DE\s+({meses}))?\s+A\s+(\d{{1,2}})\s+DE\s+({meses})\s+(?:DE\s+)?(\d{{4}})", t)
    if m:
        mes = MESES.index(m.group(4)) + 1
        return f"{mes:02d}/{m.group(5)}", m.group(0).title().replace(" De ", " de ").replace(" A ", " a ").replace("Periodo", "Período")
    # "PERIODO DE 01/02/2026 A 28/02/2026"
    m = re.search(r"PERIODO\s+(?:DE\s+)?(\d{1,2})/(\d{1,2})/(\d{4})\s+A\s+(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if m:
        return f"{int(m.group(5)):02d}/{m.group(6)}", m.group(0).capitalize().replace("Periodo", "Período")
    # "REFERENTE A FEVEREIRO/2026", "COMPETENCIA: 02/2026"
    m = re.search(rf"\b({meses})\s*(?:/|DE)\s*(\d{{4}})", t)
    if m:
        return f"{MESES.index(m.group(1)) + 1:02d}/{m.group(2)}", m.group(0).title()
    m = re.search(r"(?:REFERENTE|COMPETENCIA)\D{0,15}(\d{2})/(\d{4})", t)
    if m:
        return f"{m.group(1)}/{m.group(2)}", m.group(0).capitalize().replace("Periodo", "Período")
    return "", ""


def extrair_paciente(texto: str) -> str:
    t = re.sub(r"[ \t]+", " ", texto)
    # Formato das suas notas: "BENEFICIÁRIO PELO TRATAMENTO: NOME" (no fim da descrição)
    achou = list(re.finditer(r"BENEFICI[ÁA]RIO\s+PELO\s+TRATAMENTO\s*:\s*([^\n,;:.\d()]+)", t, flags=re.IGNORECASE))
    if achou:
        nome = re.split(r"\s+(?:CPF|CNPJ|COMPREENDIDO|REFERENTE|PER[ÍI]ODO|VALOR)\b", achou[-1].group(1),
                        flags=re.IGNORECASE)[0]
        nome = nome.strip(" .-:")
        if nome:
            return nome.upper()
    # Rótulos possíveis; o nome vai até CPF/separador/fim de linha.
    padrao = (r"(?:PACIENTE|BENEFICI[ÁA]RIO|CLIENTE\s*/?\s*PACIENTE|ATENDIMENTO\s+(?:A|AO|DO|DA)\s+PACIENTE)"
              r"\s*[:\-–]?\s*([A-Za-zÀ-ÿ'´`.]+(?:\s+(?:d[aeo]s?|e|[A-Za-zÀ-ÿ'´`.]+)){0,7})")
    for m in re.finditer(padrao, t, flags=re.IGNORECASE):
        nome = m.group(1)
        # corta em palavras que indicam fim do nome
        nome = re.split(r"\s+(?:CPF|CNPJ|COMPREENDIDO|REFERENTE|PER[ÍI]ODO|VALOR|SERVI[ÇC]O|NF|NOTA|DATA|TOTAL)\b",
                        nome, flags=re.IGNORECASE)[0]
        nome = nome.strip(" .-:")
        if len(nome.split()) >= 2 and not re.fullmatch(r"(?i)(nome|do|da|de)", nome):
            return nome.upper()
    return ""


# ---------------------------------------------------------------------- XML
def _limpa_ns(root: ET.Element) -> None:
    for el in root.iter():
        if isinstance(el.tag, str) and "}" in el.tag:
            el.tag = el.tag.split("}", 1)[1]


def _achar(root, *tags) -> str:
    """Primeiro texto não vazio de qualquer uma das tags (busca em qualquer nível)."""
    for tag in tags:
        for el in root.iter(tag):
            if el.text and el.text.strip():
                return el.text.strip()
    return ""


def _float_xml(txt: str) -> float | None:
    if not txt:
        return None
    try:
        return float(txt) if "," not in txt else num_br(txt)
    except ValueError:
        return None


def ler_xml(caminho: str) -> Nota:
    nota = Nota(arquivo=os.path.basename(caminho))
    root = ET.parse(caminho).getroot()
    _limpa_ns(root)

    nota.numero = _achar(root, "nNF", "nNFSe", "NumeroNfse", "Numero", "nDFSe").lstrip("0") or \
        _achar(root, "nNF", "nNFSe", "NumeroNfse", "Numero")

    # Bruto
    bruto = _float_xml(_achar(root, "vServ", "ValorServicos", "vServPrest", "vNF", "vProd", "ValorTotal"))
    if bruto is None:  # NFS-e nacional: <vServ> dentro de <vServPrest>
        for el in root.iter("vServPrest"):
            bruto = _float_xml(_achar(el, "vServ"))
    nota.valor_bruto = bruto

    # Líquido
    liq = _float_xml(_achar(root, "ValorLiquidoNfse", "vLiq", "vLiquido", "ValorLiquido"))
    if liq is None and bruto is not None:
        # sem campo de líquido: bruto - retenções/descontos conhecidos
        desc = 0.0
        achou = False
        for tag in ("vTotTribFed", ):  # tributos informativos não descontam; ignorados
            pass
        for tag in ("ValorIssRetido", "ValorPis", "ValorCofins", "ValorInss", "ValorIr", "ValorCsll",
                    "vRetCSLL", "vRetIRRF", "vRetCP", "vDescIncond", "ValorDescontoIncondicionado",
                    "vPISRet", "vCOFINSRet"):
            v = _float_xml(_achar(root, tag))
            if v:
                desc += v
                achou = True
        if not achou:
            nf_total = _float_xml(_achar(root, "vNF"))
            liq = nf_total if nf_total is not None else bruto
        else:
            liq = round(bruto - desc, 2)
    nota.valor_liquido = liq

    descricao = " ".join(filter(None, (
        _achar(root, "infCpl"), _achar(root, "Discriminacao"), _achar(root, "xDescServ"),
        _achar(root, "xInfComp"), _achar(root, "OutrasInformacoes"), _achar(root, "xProd"))))
    nota.paciente = extrair_paciente(descricao)
    nota.mes_ano, nota.periodo = extrair_periodo(descricao)
    return nota


# ---------------------------------------------------------------------- PDF
def _valor_apos(texto: str, rotulos: list[str], janela: int = 160) -> float | None:
    for rot in rotulos:
        for m in re.finditer(rot, texto, flags=re.IGNORECASE):
            trecho = texto[m.end(): m.end() + janela]
            v = re.search(VALOR, trecho)
            if v:
                return num_br(v.group(1))
    return None


def ler_pdf(caminho: str) -> Nota:
    import pymupdf as fitz
    nota = Nota(arquivo=os.path.basename(caminho))
    with fitz.open(caminho) as doc:
        texto = "\n".join(p.get_text() for p in doc)
    if len(texto.strip()) < 30:
        nota.aviso = "PDF sem texto (provavelmente escaneado/imagem). Use o XML."
        return nota
    plano = re.sub(r"[ \t]*\n[ \t]*", "\n", texto)

    # Número
    for pad in (r"N[º°o]?\.?\s*(?:DA\s+)?(?:NOTA|NF-?E?|NFS-?E)\s*[:\-]?\s*\n?\s*([\d\.]{1,15})\b",
                r"N[ÚU]MERO\s+(?:DA\s+)?(?:NOTA|NFS-?E|NF-?E)?\s*[:\-]?\s*\n?\s*([\d\.]{1,15})\b",
                r"\bN[º°]\s*[:\-]?\s*([\d\.]{1,15})\b",
                r"NF-?E?\s*(?:N[º°o]?)?\s*[:\-]?\s*([\d\.]{3,15})\b"):
        m = re.search(pad, plano, flags=re.IGNORECASE)
        if m:
            nota.numero = m.group(1).replace(".", "").lstrip("0") or m.group(1)
            break
    if not nota.numero:
        m = re.search(r"NOTA.{0,40}?(\d{1,9})", plano, flags=re.IGNORECASE | re.DOTALL)
        if m:
            nota.numero = m.group(1).lstrip("0")

    nota.valor_bruto = _valor_apos(plano, [
        r"VALOR\s+TOTAL\s+DOS\s+SERVI[ÇC]OS", r"VALOR\s+DOS\s+SERVI[ÇC]OS",
        r"VALOR\s+TOTAL\s+DA\s+NOTA", r"VALOR\s+TOTAL\s+DA\s+NF", r"VALOR\s+BRUTO", r"V\.?\s*TOTAL\s+DA\s+NOTA"])
    nota.valor_liquido = _valor_apos(plano, [
        r"VALOR\s+L[ÍI]QUIDO(?:\s+DA\s+NOTA|\s+DA\s+NFS-?E)?", r"L[ÍI]QUIDO\s+A\s+RECEBER", r"VALOR\s+A\s+RECEBER"])
    if nota.valor_liquido is None and nota.valor_bruto is not None:
        nota.aviso += "Valor líquido não encontrado no PDF (conferir). "

    nota.paciente = extrair_paciente(plano)
    nota.mes_ano, nota.periodo = extrair_periodo(plano)
    return nota


# ------------------------------------------------------------------ geral
def ler_arquivo(caminho: str) -> Nota:
    try:
        ext = os.path.splitext(caminho)[1].lower()
        nota = ler_xml(caminho) if ext == ".xml" else ler_pdf(caminho)
    except Exception as e:  # noqa: BLE001
        return Nota(arquivo=os.path.basename(caminho), aviso=f"Erro ao ler: {e}")
    faltando = [n for n, v in (("paciente", nota.paciente), ("número", nota.numero),
                               ("valor bruto", nota.valor_bruto), ("valor líquido", nota.valor_liquido),
                               ("mês/ano", nota.mes_ano)) if v in ("", None)]
    if faltando:
        nota.aviso += "Não encontrado: " + ", ".join(faltando)
    return nota


def coletar(caminhos: list[str]) -> list[str]:
    achados = []
    for c in caminhos:
        if os.path.isdir(c):
            for raiz, _, arqs in os.walk(c):
                achados += [os.path.join(raiz, a) for a in sorted(arqs)
                            if a.lower().endswith((".pdf", ".xml"))]
        elif c.lower().endswith((".pdf", ".xml")):
            achados.append(c)
    # Se existir PDF e XML da mesma nota (mesmo nome-base), prefere o XML.
    xmls = {os.path.splitext(a)[0].lower() for a in achados if a.lower().endswith(".xml")}
    return [a for a in achados if a.lower().endswith(".xml") or os.path.splitext(a)[0].lower() not in xmls]


def salvar_xlsx(notas: list[Nota], destino: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    cab = ["Paciente", "Nº da NF", "Valor bruto", "Valor líquido", "Mês/Ano", "Período", "Arquivo", "Avisos"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
        c.alignment = Alignment(horizontal="center")
    for n in notas:
        ws.append([n.paciente, int(n.numero) if n.numero.isdigit() else n.numero,
                   n.valor_bruto, n.valor_liquido, n.mes_ano, n.periodo, n.arquivo, n.aviso])
        if n.aviso:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(ws.max_row, 5).number_format = "@"
    for linha in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in linha:
            c.number_format = 'R$ #,##0.00'
    if len(notas):
        ultima = ws.max_row + 1
        ws.cell(ultima, 1, "TOTAL").font = Font(bold=True)
        for col, letra in ((3, "C"), (4, "D")):
            c = ws.cell(ultima, col, f"=SUM({letra}2:{letra}{ultima - 1})")
            c.font = Font(bold=True)
            c.number_format = 'R$ #,##0.00'
    for col, larg in zip("ABCDEFGH", (38, 10, 14, 14, 10, 34, 34, 40)):
        ws.column_dimensions[col].width = larg
    ws.freeze_panes = "A2"
    wb.save(destino)


# -------------------------------------------------------------------- GUI
def abrir_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    notas: list[Nota] = []
    root = tk.Tk()
    root.title("Leitor de Notas Fiscais -> Planilha")
    root.geometry("1100x520")

    topo = ttk.Frame(root, padding=8)
    topo.pack(fill="x")
    colunas = ("paciente", "numero", "bruto", "liquido", "mes", "arquivo", "aviso")
    tree = ttk.Treeview(root, columns=colunas, show="headings")
    for c, t, w in zip(colunas, ("Paciente", "Nº NF", "Bruto", "Líquido", "Mês/Ano", "Arquivo", "Avisos"),
                       (260, 70, 90, 90, 80, 220, 260)):
        tree.heading(c, text=t)
        tree.column(c, width=w)
    tree.tag_configure("aviso", background="#FFF2CC")
    tree.pack(fill="both", expand=True, padx=8)
    status = ttk.Label(root, text="Escolha arquivos PDF/XML ou uma pasta.", padding=8)
    status.pack(fill="x")

    def fmt(v):
        return "" if v is None else f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def processar(caminhos):
        for arq in coletar(list(caminhos)):
            n = ler_arquivo(arq)
            notas.append(n)
            tree.insert("", "end", tags=("aviso",) if n.aviso else (),
                        values=(n.paciente, n.numero, fmt(n.valor_bruto), fmt(n.valor_liquido),
                                n.mes_ano, n.arquivo, n.aviso))
        status.config(text=f"{len(notas)} nota(s) carregada(s). Linhas amarelas precisam de conferência.")

    def add_arquivos():
        processar(filedialog.askopenfilenames(filetypes=[("Notas", "*.pdf *.xml")]))

    def add_pasta():
        p = filedialog.askdirectory()
        if p:
            processar([p])

    def limpar():
        notas.clear()
        tree.delete(*tree.get_children())
        status.config(text="Lista limpa.")

    def salvar():
        if not notas:
            messagebox.showinfo("Nada a salvar", "Carregue alguma nota primeiro.")
            return
        d = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="notas.xlsx",
                                         filetypes=[("Excel", "*.xlsx")])
        if d:
            try:
                salvar_xlsx(notas, d)
            except PermissionError:
                messagebox.showerror("Erro", "Feche a planilha no Excel e tente de novo.")
                return
            messagebox.showinfo("Pronto", f"Planilha salva em:\n{d}")
            os.startfile(d) if sys.platform == "win32" else None

    for txt, cmd in (("Adicionar arquivos…", add_arquivos), ("Adicionar pasta…", add_pasta),
                     ("Limpar", limpar), ("Salvar planilha…", salvar)):
        ttk.Button(topo, text=txt, command=cmd).pack(side="left", padx=4)
    root.mainloop()


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        abrir_gui()
    else:
        resultado = [ler_arquivo(a) for a in coletar(args)]
        for r in resultado:
            print(asdict(r))
        salvar_xlsx(resultado, "notas.xlsx")
        print("Salvo em notas.xlsx")
